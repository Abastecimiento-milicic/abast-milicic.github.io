# -*- coding: utf-8 -*-
import os
import sys
import glob
import re
import datetime
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def apply_excel_table_styling(ws):
    """
    Aplica estilo profesional corporativo a la hoja Excel:
    - Encabezado con azul corporativo (#1B365D), bordes y texto en blanco.
    - Encabezado destacado para 'Caract.planif.nec.' (#0284C7) y 'Planificado' (#047857).
    - Filas alternadas suaves (zebra striping: blanco y #F8FAFC).
    - Bordes suaves en todas las celdas para mantener la cuadrícula nítida.
    - Destacado con color en la columna 'Caract.planif.nec.':
        * Planificados (V1, RE, CL, etc.): Fondo verde menta (#DCFCE7), texto verde oscuro (#166534), negrita.
        * No planificados (NO PLANIFICADO, etc.): Fondo rojo suave (#FEE2E2), texto rojo oscuro (#991B1B), negrita.
    - Destacado con color en la columna 'Planificado':
        * SÍ: Fondo verde menta (#DCFCE7), texto verde oscuro (#166534), negrita.
        * NO: Fondo rojo suave (#FEE2E2), texto rojo oscuro (#991B1B), negrita.
    - Formato numérico para cantidades (#,##0) y valores monetarios ($ #,##0.00).
    - Fila de encabezado congelada (Freeze panes en A2).
    - Auto-filtro activado en toda la tabla.
    - Auto-ajuste de ancho de columnas.
    """
    max_row = ws.max_row
    max_col = ws.max_column
    if max_row < 1 or max_col < 1:
        return

    # Paletas de color
    header_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_sky = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")      # Caract.planif.nec.
    header_emerald = PatternFill(start_color="047857", end_color="047857", fill_type="solid")  # Planificado
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    border_cell = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    border_header = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='medium', color='0F172A'),
        bottom=Side(style='medium', color='0F172A')
    )

    zebra_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    zebra_odd = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name="Segoe UI", size=10, bold=True, color="166534")

    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    default_font = Font(name="Segoe UI", size=10, color="1F2937")

    headers = [str(ws.cell(row=1, column=c).value or '').strip() for c in range(1, max_col + 1)]

    caract_col = None
    planif_col = None
    for c_idx, h in enumerate(headers, 1):
        if h == 'Caract.planif.nec.':
            caract_col = c_idx
        elif h == 'Planificado':
            planif_col = c_idx

    currency_cols = {
        'Valor libre util.', 'Val.trans.c/cond.', 'Valor en insp.cal.',
        'Valor no libre', 'Valor stock bloq.', 'Val.stock bl.dev.'
    }
    qty_cols = {
        'Punto de pedido', 'Stock alm.máx.', 'Libre utilización',
        'Trans./Trasl.', 'Inspecc.de calidad', 'Stock no libre',
        'Bloqueado', 'Devoluciones'
    }
    center_cols = {
        'Centro', 'Almacén', 'Planificado', 'Caract.planif.nec.',
        'Unidad medida base', 'Moneda', 'Tipo material'
    }

    # 1. Encabezados
    ws.row_dimensions[1].height = 28
    for c in range(1, max_col + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.border = border_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if c == caract_col:
            cell.fill = header_sky
        elif c == planif_col:
            cell.fill = header_emerald
        else:
            cell.fill = header_navy

    # 2. Filas de datos
    for r in range(2, max_row + 1):
        ws.row_dimensions[r].height = 20
        is_odd = (r % 2 == 1)
        row_base_fill = zebra_odd if is_odd else zebra_even

        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            h_name = headers[c - 1]
            val = cell.value

            cell.border = border_cell
            cell.font = default_font
            cell.fill = row_base_fill

            if h_name in center_cols:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif h_name in currency_cols or h_name in qty_cols:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if val is not None and isinstance(val, (int, float)):
                if h_name in currency_cols:
                    cell.number_format = '$ #,##0.00'
                elif h_name in qty_cols:
                    cell.number_format = '#,##0'

            # Destacado con color en Caract.planif.nec.
            if c == caract_col:
                val_str = str(val or '').strip().upper()
                if val_str and val_str not in ('NO PLANIFICADO', 'ND', '-', 'SIN PLANIFICAR', 'NONE'):
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = red_fill
                    cell.font = red_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Destacado con color en Planificado
            elif c == planif_col:
                val_str = str(val or '').strip().upper()
                if val_str in ('SÍ', 'SI', 'YES', '1'):
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = red_fill
                    cell.font = red_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Congelar encabezado
    ws.freeze_panes = 'A2'

    # 4. Auto-filtro
    ws.auto_filter.ref = ws.dimensions

    # 5. Líneas de cuadrícula
    try:
        if ws.views.sheetView:
            ws.views.sheetView[0].showGridLines = True
    except Exception:
        pass

    # 6. Auto-ancho de columnas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v = cell.value
            if v is not None:
                s = str(v)
                if len(s) > max_len:
                    max_len = len(s)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 11), 45)

def norm_alm(v):
    if pd.isna(v):
        return ''
    s = str(v).strip()
    s = re.sub(r'\.0$', '', s)
    return s.upper()

def alm_keys(alm_str):
    k = norm_alm(alm_str)
    if not k:
        return []
    res = [k]
    k_no0 = k.lstrip('0')
    if k_no0 and k_no0 != k:
        res.append(k_no0)
    # Si tiene sufijo como 378B, probar con base 378
    m = re.match(r'^(\d+)[A-Z\-]?$', k_no0)
    if m:
        base = m.group(1)
        if base not in res:
            res.append(base)
    return res

def norm_mat(v):
    if pd.isna(v):
        return ''
    s = str(v).strip()
    s = re.sub(r'\.0$', '', s)
    return s.lstrip('0')

def load_analisis_mm(csv_path):
    encodings = ['utf-8-sig', 'utf-8', 'latin1', 'cp1252']
    df_mm = None
    for enc in encodings:
        try:
            df_mm = pd.read_csv(csv_path, sep=';', encoding=enc, low_memory=False)
            break
        except Exception:
            continue
    if df_mm is None:
        raise RuntimeError(f"No se pudo leer {csv_path} con las codificaciones habituales.")
    return df_mm

def load_auxiliary_descriptions(base_csv_path):
    desc_map = {}
    if os.path.exists(base_csv_path):
        try:
            df_aux = pd.read_csv(base_csv_path, sep=';', usecols=['CODIGO ITEM', 'DESCRIPCION ITEM'], encoding='utf-8-sig', low_memory=False)
            df_aux = df_aux.dropna(subset=['CODIGO ITEM', 'DESCRIPCION ITEM'])
            for _, r in df_aux.iterrows():
                m = norm_mat(r['CODIGO ITEM'])
                if m and m not in desc_map:
                    desc_map[m] = str(r['DESCRIPCION ITEM']).strip()
        except Exception:
            pass
    return desc_map

def load_mm60_metadata(data_dir):
    """
    Carga el catálogo maestro desde MM60 (ej: MM60 06-10-2026.xlsx) para obtener
    'Tipo material', 'Texto breve de material' y 'Grupo de compras'.
    """
    pattern_xlsx = os.path.join(data_dir, "*MM60*.xlsx")
    pattern_xls = os.path.join(data_dir, "*MM60*.XLSX")
    candidates = glob.glob(pattern_xlsx) + glob.glob(pattern_xls)
    candidates = [f for f in candidates if not os.path.basename(f).startswith('~$')]
    if not candidates:
        print("[AVISO] No se encontró archivo de MM60 en", data_dir)
        return {}

    candidates.sort(key=os.path.getmtime, reverse=True)
    fpath = candidates[0]
    fname = os.path.basename(fpath)
    print(f"Cargando catálogo maestro desde {fname}...")
    try:
        cols = ['Material', 'Tipo material', 'Texto breve de material', 'Grupo de compras']
        df = pd.read_excel(fpath, usecols=lambda c: c in cols)
        df = df.dropna(subset=['Material']).drop_duplicates(subset=['Material']).copy()
        df['mat_norm'] = df['Material'].apply(norm_mat)
        cols_present = [c for c in ['Tipo material', 'Texto breve de material', 'Grupo de compras'] if c in df.columns]
        mm60_map = df.set_index('mat_norm')[cols_present].to_dict(orient='index')
        print(f"-> {len(mm60_map):,d} materiales indexados desde {fname}.")
        return mm60_map
    except Exception as e:
        print(f"[ERROR al cargar {fname}]: {e}")
        return {}

def run_combination(base_dir=None):
    if base_dir is None:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

    data_dir = os.path.join(base_dir, 'data')
    mb52_dir = os.path.join(data_dir, 'Almacen MB52')
    mm_path = os.path.join(data_dir, 'ANALISIS-MM.csv')
    aux_path = os.path.join(data_dir, 'Base_Final_365_372.csv')
    out_dir = os.path.join(mb52_dir, 'Combinados')

    os.makedirs(out_dir, exist_ok=True)

    print("=" * 75)
    print("      COMBINACIÓN DE STOCK ALMACÉN MB52 CON ANÁLISIS MM (SAP)")
    print("=" * 75)
    print(f"Directorio de datos : {data_dir}")
    print(f"Carpeta MB52        : {mb52_dir}")
    print(f"Carpeta de salida   : {out_dir}")
    print("-" * 75)

    if not os.path.exists(mm_path):
        print(f"[ERROR] No se encontró el archivo {mm_path}")
        return

    print("Cargando maestro de planificación ANALISIS-MM.csv...")
    df_mm = load_analisis_mm(mm_path)
    print(f"-> {len(df_mm):,d} registros cargados de ANALISIS-MM.")

    print("Indexando claves de planificación por Almacén y Material...")
    mm_exact = {}
    mm_mat_meta = {}

    for _, row in df_mm.iterrows():
        mat = norm_mat(row.get('Material'))
        if not mat:
            continue

        for a in alm_keys(row.get('Almacén', '')):
            mm_exact[(a, mat)] = row

        if mat not in mm_mat_meta:
            mm_mat_meta[mat] = {
                'Texto breve de material': str(row.get('Texto breve de material', '') or '').strip(),
                'Rubro': str(row.get('Rubro', '') or '').strip(),
                'SubRubro': str(row.get('SubRubro', '') or '').strip(),
                'Tipo material': str(row.get('Tipo material', '') or '').strip(),
                'Grupo de compras': str(row.get('Grupo de compras', '') or '').strip()
            }

    # Cargar descripciones auxiliares si existen
    aux_descs = load_auxiliary_descriptions(aux_path)
    if aux_descs:
        print(f"-> {len(aux_descs):,d} descripciones adicionales disponibles desde Base_Final.")

    # Cargar catálogo maestro MM60 (Tipo material, Texto breve, Grupo de compras)
    mm60_map = load_mm60_metadata(data_dir)

    # Buscar archivos MB52
    mb52_pattern = os.path.join(mb52_dir, "*.XLS*")
    all_files = glob.glob(mb52_pattern) + glob.glob(os.path.join(mb52_dir, "*.xls*"))
    all_files = sorted(list(set(all_files)))

    # Filtrar para no reprocesar archivos de la carpeta Combinados
    raw_files = [f for f in all_files if 'combinado' not in os.path.basename(f).lower() and not f.startswith(out_dir)]

    if not raw_files:
        print(f"[AVISO] No se encontraron archivos de MB52 para procesar en {mb52_dir}")
        return

    print(f"Se encontraron {len(raw_files)} archivos de MB52 para procesar.\n")

    summary_list = []
    all_combined_dfs = []

    for fpath in raw_files:
        fname = os.path.basename(fpath)
        base_name_no_ext = os.path.splitext(fname)[0]
        print(f"Procesando: {fname} ...")

        try:
            df_mb = pd.read_excel(fpath)
        except Exception as e:
            print(f"  [ERROR] No se pudo leer {fname}: {e}")
            continue

        if 'Material' not in df_mb.columns:
            print(f"  [SKIP] No tiene columna 'Material'. Omitiendo.")
            continue

        # Eliminar filas vacías o fila de total resumen
        df_clean = df_mb.dropna(subset=['Material']).copy()
        if df_clean.empty:
            print(f"  [SKIP] Archivo vacío después de filtrar materiales.")
            continue

        # Enriquecer cada fila
        enriched_rows = []
        for _, r in df_clean.iterrows():
            mat_raw = r['Material']
            mat_clean = norm_mat(mat_raw)
            alm_raw = r.get('Almacén', '')
            alm_candidates = alm_keys(alm_raw)

            row_mm = None
            for a in alm_candidates:
                if (a, mat_clean) in mm_exact:
                    row_mm = mm_exact[(a, mat_clean)]
                    break

            tipo_mat = ''
            if row_mm is not None:
                caract = str(row_mm.get('Caract.planif.nec.', '') or '').strip() or 'PLANIFICADO'
                planif = 'SI'
                estado_mm = str(row_mm.get('Estado', '') or '').strip() or 'PLANIFICADO'
                pp = row_mm.get('Punto de pedido', 0)
                stk_max = row_mm.get('Stock alm.máx.', 0)
                desc = str(row_mm.get('Texto breve de material', '') or '').strip()
                rubro = str(row_mm.get('Rubro', '') or '').strip()
                subrubro = str(row_mm.get('SubRubro', '') or '').strip()
                grp_comp = str(row_mm.get('Grupo de compras', '') or '').strip()
                tipo_mat = str(row_mm.get('Tipo material', '') or '').strip()
            else:
                caract = 'NO PLANIFICADO'
                planif = 'NO'
                estado_mm = 'NO PLANIFICADO'
                pp = 0
                stk_max = 0
                # Búsqueda de metadata por código de material
                meta = mm_mat_meta.get(mat_clean, {})
                desc = meta.get('Texto breve de material', '')
                rubro = meta.get('Rubro', '')
                subrubro = meta.get('SubRubro', '')
                grp_comp = meta.get('Grupo de compras', '')
                tipo_mat = meta.get('Tipo material', '')

            # Complementar con catálogo maestro MM60
            mm60_info = mm60_map.get(mat_clean, {})
            if not tipo_mat:
                tipo_mat = str(mm60_info.get('Tipo material', '') or '').strip()
            if not desc or desc == 'SIN DESCRIPCIÓN':
                mm60_desc = str(mm60_info.get('Texto breve de material', '') or '').strip()
                if mm60_desc:
                    desc = mm60_desc
            if not grp_comp:
                mm60_grp = str(mm60_info.get('Grupo de compras', '') or '').strip()
                if mm60_grp:
                    grp_comp = mm60_grp

            if not desc and mat_clean in aux_descs:
                desc = aux_descs[mat_clean]

            # Armar registro enriquecido
            rec = {
                'Material': int(float(mat_clean)) if mat_clean.isdigit() else mat_clean,
                'Texto breve de material': desc or 'SIN DESCRIPCIÓN',
                'Tipo material': tipo_mat or '',
                'Centro': str(r.get('Centro', '') or '').strip(),
                'Almacén': norm_alm(r.get('Almacén', '')) or str(r.get('Almacén', '') or '').strip(),
                'Planificado': planif,
                'Caract.planif.nec.': caract,
                'Estado MM': estado_mm,
                'Punto de pedido': pp,
                'Stock alm.máx.': stk_max,
                'Libre utilización': r.get('Libre utilización', 0),
                'Unidad medida base': r.get('Unidad medida base', ''),
                'Valor libre util.': r.get('Valor libre util.', 0.0),
                'Moneda': r.get('Moneda', 'ARS'),
                'Rubro': rubro or 'SIN RUBRO',
                'SubRubro': subrubro or '',
                'Grupo de compras': grp_comp or '',
                'Trans./Trasl.': r.get('Trans./Trasl.', 0),
                'Val.trans.c/cond.': r.get('Val.trans.c/cond.', 0.0),
                'Inspecc.de calidad': r.get('Inspecc.de calidad', 0),
                'Valor en insp.cal.': r.get('Valor en insp.cal.', 0.0),
                'Stock no libre': r.get('Stock no libre', 0),
                'Valor no libre': r.get('Valor no libre', 0.0),
                'Bloqueado': r.get('Bloqueado', 0),
                'Valor stock bloq.': r.get('Valor stock bloq.', 0.0),
                'Devoluciones': r.get('Devoluciones', 0),
                'Val.stock bl.dev.': r.get('Val.stock bl.dev.', 0.0)
            }
            enriched_rows.append(rec)

        df_out = pd.DataFrame(enriched_rows)

        # Guardar archivo individual en Combinados
        out_excel = os.path.join(out_dir, f"{base_name_no_ext} - Combinado.xlsx")
        try:
            with pd.ExcelWriter(out_excel, engine='openpyxl') as writer:
                df_out.to_excel(writer, sheet_name='Stock_Planificacion', index=False)
                apply_excel_table_styling(writer.sheets['Stock_Planificacion'])
            print(f"  -> Guardado con formato: {out_excel}")
        except Exception as e:
            print(f"  [ERROR al guardar Excel]: {e}")

        # Acumular para consolidado
        df_out_conso = df_out.copy()
        df_out_conso.insert(0, 'Archivo_Origen', fname)
        all_combined_dfs.append(df_out_conso)

        # Métricas de resumen
        tot = len(df_out)
        plan_c = (df_out['Planificado'] == 'SI').sum()
        no_plan_c = tot - plan_c
        pct_p = (plan_c / tot * 100) if tot else 0
        v_plan = df_out.loc[df_out['Planificado'] == 'SI', 'Valor libre util.'].fillna(0).sum()
        v_no_plan = df_out.loc[df_out['Planificado'] == 'NO', 'Valor libre util.'].fillna(0).sum()

        summary_list.append({
            'Archivo': fname,
            'Total_Items': tot,
            'Planificados': plan_c,
            '%_Planif': round(pct_p, 1),
            'No_Planificados': no_plan_c,
            'Valor_Planificado': v_plan,
            'Valor_No_Planificado': v_no_plan
        })

    # Guardar Consolidado Maestro
    if all_combined_dfs:
        print("\nGenerando Consolidado Maestro de todos los almacenes...")
        df_all = pd.concat(all_combined_dfs, ignore_index=True)
        
        conso_xlsx = os.path.join(out_dir, "MB52_CONSOLIDADO_COMBINADO.xlsx")
        conso_csv = os.path.join(out_dir, "MB52_CONSOLIDADO_COMBINADO.csv")

        try:
            with pd.ExcelWriter(conso_xlsx, engine='openpyxl') as writer:
                df_all.to_excel(writer, sheet_name='Consolidado_MB52', index=False)
                apply_excel_table_styling(writer.sheets['Consolidado_MB52'])
            print(f"-> Consolidado Excel con formato : {conso_xlsx}")
        except Exception as e:
            print(f"[ERROR al guardar Consolidado Excel]: {e}")

        try:
            df_all.to_csv(conso_csv, sep=';', index=False, encoding='utf-8-sig')
            print(f"-> Consolidado CSV   : {conso_csv}")
        except Exception as e:
            print(f"[ERROR al guardar Consolidado CSV]: {e}")

    # Imprimir tabla de resumen
    if summary_list:
        df_resumen = pd.DataFrame(summary_list)
        print("\n" + "=" * 95)
        print(f"{'RESUMEN FINAL DE COMBINACIÓN':^95}")
        print("=" * 95)
        print(f"{'Archivo':<22} | {'Items':>6} | {'Planif':>7} | {'% Plan':>7} | {'No Plan':>7} | {'Valor Planif (ARS)':>18} | {'Valor No Plan (ARS)':>18}")
        print("-" * 95)
        
        tot_items = df_resumen['Total_Items'].sum()
        tot_plan = df_resumen['Planificados'].sum()
        tot_no_plan = df_resumen['No_Planificados'].sum()
        tot_val_p = df_resumen['Valor_Planificado'].sum()
        tot_val_np = df_resumen['Valor_No_Planificado'].sum()
        tot_pct = (tot_plan / tot_items * 100) if tot_items else 0

        for _, r in df_resumen.iterrows():
            print(f"{r['Archivo']:<22} | {r['Total_Items']:6,d} | {r['Planificados']:7,d} | {r['%_Planif']:6.1f}% | {r['No_Planificados']:7,d} | ${r['Valor_Planificado']:17,.2f} | ${r['Valor_No_Planificado']:17,.2f}")
        print("-" * 95)
        print(f"{'TOTAL GENERAL':<22} | {tot_items:6,d} | {tot_plan:7,d} | {tot_pct:6.1f}% | {tot_no_plan:7,d} | ${tot_val_p:17,.2f} | ${tot_val_np:17,.2f}")
        print("=" * 95)

if __name__ == '__main__':
    run_combination()
